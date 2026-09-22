"""R4 raw transaction standardization, source-order recovery and reconciliation.

This module receives an R3-confirmed mapping.  It preserves source direction,
signed source amount, sheet and row location; it neither aggregates dates nor
writes files or databases.  Calendar, daily aggregation and persistence remain
separate later responsibilities.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
import hashlib
from pathlib import Path
import re
from typing import Iterable

from openpyxl import load_workbook

from runtime_field_mapping import ConfirmedFieldMapping, inspect_workbook_headers


CODE_NORMALIZED = "TN_NORMALIZED_001"
CODE_NORMALIZED_WITH_INVALID_ROWS = "TN_NORMALIZED_WITH_INVALID_ROWS_001"
CODE_MAPPING_FINGERPRINT_MISMATCH = "TN_MAPPING_FINGERPRINT_MISMATCH_001"
CODE_DIRECTION_UNKNOWN = "TN_DIRECTION_UNKNOWN_001"
CODE_REQUIRED_VALUE_MISSING = "TN_REQUIRED_VALUE_MISSING_001"
CODE_VALUE_UNPARSEABLE = "TN_VALUE_UNPARSEABLE_001"

BALANCE_TOLERANCE_CNY = Decimal("0.01")
# The frozen order rule gives source row order priority.  Excel serial values
# can retain milliseconds that the bank's visible export rounds to seconds, so
# a sub-second timestamp reversal is treated as a source-order tie, not as an
# order-quality failure.  Larger reversals remain auditable warnings.
BOOKING_TIME_ORDER_TOLERANCE = timedelta(seconds=1)


@dataclass(frozen=True)
class RawTransaction:
    source_row_number: int
    source_sheet: str
    booking_datetime: datetime
    source_direction: str
    source_amount_cny: Decimal
    post_transaction_balance_cny: Decimal


@dataclass(frozen=True)
class NormalizedTransaction:
    transaction_id: str
    account_token: str
    source_file_sha256: str
    source_sheet: str
    source_row_number: int
    normalized_sequence: int
    booking_datetime: datetime
    booking_date: str
    source_direction: str
    source_amount_cny: Decimal
    effective_direction: str
    debit_amount_cny: Decimal
    credit_amount_cny: Decimal
    post_transaction_balance_cny: Decimal
    normalization_reason: str | None


@dataclass(frozen=True)
class InvalidTransactionRow:
    source_row_number: int
    source_sheet: str
    code: str
    booking_date: str | None


@dataclass(frozen=True)
class ReconciliationFailure:
    previous_source_row_number: int
    source_row_number: int
    booking_date: str


@dataclass(frozen=True)
class ReconciliationSummary:
    checks: int
    failures: int
    failure_rate: Decimal | None
    failure_examples: tuple[ReconciliationFailure, ...]
    skipped_due_to_invalid_rows: int = 0


@dataclass(frozen=True)
class NormalizationResult:
    code: str
    message_zh: str
    input_transaction_rows: int
    normalized_transaction_rows: int
    transactions: tuple[NormalizedTransaction, ...]
    invalid_rows: tuple[InvalidTransactionRow, ...]
    chosen_order: str | None
    booking_datetime_regression_rows: tuple[int, ...]
    reconciliation: ReconciliationSummary
    write_count: int = 0


def normalize_workbook(
    path: Path,
    confirmed_mapping: ConfirmedFieldMapping,
    *,
    account_token: str,
) -> NormalizationResult:
    """Read a confirmed raw workbook and return in-memory normalized records."""

    proposal = inspect_workbook_headers(path)
    if (
        proposal.header_fingerprint != confirmed_mapping.header_fingerprint
        or proposal.input_mode != "raw_transaction_full_snapshot_xlsx"
        or confirmed_mapping.input_mode != "raw_transaction_full_snapshot_xlsx"
        or proposal.field_to_column != confirmed_mapping.field_to_column
        or proposal.sheet_name != confirmed_mapping.sheet_name
        or proposal.header_row_number != confirmed_mapping.header_row_number
    ):
        return _blocked_mapping_result()

    source_hash = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    raw_rows, invalid_rows, input_count = _read_raw_rows(Path(path), confirmed_mapping)
    normalized = normalize_raw_transactions(
        raw_rows,
        account_token=account_token,
        source_file_sha256=source_hash,
        known_invalid_rows=invalid_rows,
    )
    combined_invalid = tuple(invalid_rows) + normalized.invalid_rows
    if combined_invalid:
        return NormalizationResult(
            code=CODE_NORMALIZED_WITH_INVALID_ROWS,
            message_zh="部分逐笔记录无法标准化，已保留匿名来源定位；本步骤不写数据库。",
            input_transaction_rows=input_count,
            normalized_transaction_rows=normalized.normalized_transaction_rows,
            transactions=normalized.transactions,
            invalid_rows=combined_invalid,
            chosen_order=normalized.chosen_order,
            booking_datetime_regression_rows=normalized.booking_datetime_regression_rows,
            reconciliation=normalized.reconciliation,
        )
    return NormalizationResult(
        code=normalized.code,
        message_zh=normalized.message_zh,
        input_transaction_rows=input_count,
        normalized_transaction_rows=normalized.normalized_transaction_rows,
        transactions=normalized.transactions,
        invalid_rows=(),
        chosen_order=normalized.chosen_order,
        booking_datetime_regression_rows=normalized.booking_datetime_regression_rows,
        reconciliation=normalized.reconciliation,
    )


def normalize_raw_transactions(
    raw_transactions: Iterable[RawTransaction],
    *,
    account_token: str,
    source_file_sha256: str,
    known_invalid_rows: Iterable[InvalidTransactionRow] = (),
) -> NormalizationResult:
    """Standardize supplied raw rows and reconcile their recovered source order."""

    raw_rows = tuple(raw_transactions)
    valid_rows: list[RawTransaction] = []
    invalid_rows: list[InvalidTransactionRow] = list(known_invalid_rows)
    for row in raw_rows:
        if _effective_direction(row.source_direction, row.source_amount_cny) is None:
            invalid_rows.append(
                InvalidTransactionRow(
                    source_row_number=row.source_row_number,
                    source_sheet=row.source_sheet,
                    code=CODE_DIRECTION_UNKNOWN,
                    booking_date=row.booking_datetime.date().isoformat(),
                )
            )
        else:
            valid_rows.append(row)

    chosen_order, ordered_rows = _recover_source_order(tuple(valid_rows))
    identity_occurrences: dict[str, int] = {}
    normalized_rows: list[NormalizedTransaction] = []
    for index, row in enumerate(ordered_rows, start=1):
        identity_fingerprint = _transaction_identity_fingerprint(row, account_token=account_token)
        occurrence = identity_occurrences.get(identity_fingerprint, 0) + 1
        identity_occurrences[identity_fingerprint] = occurrence
        normalized_rows.append(
            _normalize_one(
                row,
                account_token=account_token,
                source_file_sha256=source_file_sha256,
                sequence=index,
                identity_fingerprint=identity_fingerprint,
                identity_occurrence=occurrence,
            )
        )
    transactions = tuple(normalized_rows)
    reconciliation = _reconcile(transactions, invalid_rows=tuple(invalid_rows))
    regressions = tuple(
        current.source_row_number
        for previous, current in zip(transactions, transactions[1:])
        if current.booking_datetime + BOOKING_TIME_ORDER_TOLERANCE < previous.booking_datetime
    )
    code = CODE_NORMALIZED_WITH_INVALID_ROWS if invalid_rows else CODE_NORMALIZED
    message = (
        "逐笔记录已标准化并完成相邻余额勾稽；本步骤不汇总日度数据，也不写数据库。"
        if not invalid_rows
        else "存在方向无法识别的逐笔记录，已保留匿名来源定位；本步骤不写数据库。"
    )
    return NormalizationResult(
        code=code,
        message_zh=message,
        input_transaction_rows=len(raw_rows),
        normalized_transaction_rows=len(transactions),
        transactions=transactions,
        invalid_rows=tuple(invalid_rows),
        chosen_order=chosen_order,
        booking_datetime_regression_rows=regressions,
        reconciliation=reconciliation,
    )


def _read_raw_rows(
    path: Path,
    mapping: ConfirmedFieldMapping,
) -> tuple[tuple[RawTransaction, ...], tuple[InvalidTransactionRow, ...], int]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[mapping.sheet_name]
        header_values = next(sheet.iter_rows(min_row=mapping.header_row_number, max_row=mapping.header_row_number, values_only=True))
        headers = ["" if value is None else str(value).strip() for value in header_values]
        positions = {field: headers.index(source_header) for field, source_header in mapping.field_to_column.items()}
        raw_rows: list[RawTransaction] = []
        invalid_rows: list[InvalidTransactionRow] = []
        input_count = 0
        for row_number, values in enumerate(sheet.iter_rows(min_row=mapping.header_row_number + 1, values_only=True), start=mapping.header_row_number + 1):
            core_values = {field: values[index] if index < len(values) else None for field, index in positions.items()}
            if all(_is_blank(value) for value in core_values.values()):
                continue
            input_count += 1
            try:
                booking_datetime = _parse_datetime(core_values["booking_datetime_or_booking_date"])
                source_direction = _parse_direction(core_values["direction"])
                source_amount = _parse_decimal(core_values["amount"])
                balance = _parse_decimal(core_values["post_transaction_balance"])
            except _RequiredValueMissing:
                invalid_rows.append(InvalidTransactionRow(row_number, sheet.title, CODE_REQUIRED_VALUE_MISSING, None))
                continue
            except (TypeError, ValueError, InvalidOperation):
                invalid_rows.append(InvalidTransactionRow(row_number, sheet.title, CODE_VALUE_UNPARSEABLE, _best_effort_date(core_values["booking_datetime_or_booking_date"])))
                continue
            raw_rows.append(RawTransaction(row_number, sheet.title, booking_datetime, source_direction, source_amount, balance))
        return tuple(raw_rows), tuple(invalid_rows), input_count
    finally:
        workbook.close()


def _recover_source_order(rows: tuple[RawTransaction, ...]) -> tuple[str | None, tuple[RawTransaction, ...]]:
    if not rows:
        return None, ()
    increasing = 0
    decreasing = 0
    for previous, current in zip(rows, rows[1:]):
        if current.booking_datetime > previous.booking_datetime:
            increasing += 1
        elif current.booking_datetime < previous.booking_datetime:
            decreasing += 1
    if decreasing > increasing:
        return "reverse_source_order", tuple(reversed(rows))
    return "source_order", rows


def _normalize_one(
    row: RawTransaction,
    *,
    account_token: str,
    source_file_sha256: str,
    sequence: int,
    identity_fingerprint: str,
    identity_occurrence: int,
) -> NormalizedTransaction:
    effective_direction = _effective_direction(row.source_direction, row.source_amount_cny)
    if effective_direction is None:
        raise ValueError("调用方必须先验证交易方向")
    amount = abs(row.source_amount_cny)
    is_reversal = row.source_amount_cny < 0
    debit = amount if effective_direction == "debit" else Decimal("0")
    credit = amount if effective_direction == "credit" else Decimal("0")
    identifier_text = f"r4_transaction_identity_v2|{identity_fingerprint}|{identity_occurrence}"
    transaction_id = hashlib.sha256(identifier_text.encode("utf-8")).hexdigest()
    return NormalizedTransaction(
        transaction_id=transaction_id,
        account_token=account_token,
        source_file_sha256=source_file_sha256,
        source_sheet=row.source_sheet,
        source_row_number=row.source_row_number,
        normalized_sequence=sequence,
        booking_datetime=row.booking_datetime,
        booking_date=row.booking_datetime.date().isoformat(),
        source_direction=row.source_direction,
        source_amount_cny=row.source_amount_cny,
        effective_direction=effective_direction,
        debit_amount_cny=debit,
        credit_amount_cny=credit,
        post_transaction_balance_cny=row.post_transaction_balance_cny,
        normalization_reason="negative_amount_reversal" if is_reversal else None,
    )


def _reconcile(
    transactions: tuple[NormalizedTransaction, ...],
    *,
    invalid_rows: tuple[InvalidTransactionRow, ...],
) -> ReconciliationSummary:
    checks = 0
    skipped = 0
    failures: list[ReconciliationFailure] = []
    for previous, current in zip(transactions, transactions[1:]):
        if _has_invalid_row_between(previous, current, invalid_rows):
            skipped += 1
            continue
        checks += 1
        expected_balance = previous.post_transaction_balance_cny - current.debit_amount_cny + current.credit_amount_cny
        if abs(current.post_transaction_balance_cny - expected_balance) > BALANCE_TOLERANCE_CNY:
            failures.append(
                ReconciliationFailure(
                    previous_source_row_number=previous.source_row_number,
                    source_row_number=current.source_row_number,
                    booking_date=current.booking_date,
                )
            )
    rate = (Decimal(len(failures)) / Decimal(checks)) if checks else None
    return ReconciliationSummary(checks, len(failures), rate, tuple(failures), skipped)


def _has_invalid_row_between(
    previous: NormalizedTransaction,
    current: NormalizedTransaction,
    invalid_rows: tuple[InvalidTransactionRow, ...],
) -> bool:
    if previous.source_sheet != current.source_sheet:
        return True
    lower, upper = sorted((previous.source_row_number, current.source_row_number))
    return any(
        invalid.source_sheet == previous.source_sheet and lower < invalid.source_row_number < upper
        for invalid in invalid_rows
    )


def _transaction_identity_fingerprint(row: RawTransaction, *, account_token: str) -> str:
    """Return a snapshot-independent identity basis for a parsed transaction.

    Full history exports can add rows and therefore change both their file hash
    and source row positions.  Those are source-audit attributes, not stable
    transaction identity.  Exact repeated records are separated later by their
    chronological occurrence number.
    """

    identity_parts = (
        "r4_transaction_identity_basis_v2",
        account_token,
        row.booking_datetime.isoformat(timespec="microseconds"),
        row.source_direction.strip(),
        format(row.source_amount_cny, "f"),
        format(row.post_transaction_balance_cny, "f"),
    )
    return hashlib.sha256("\x1f".join(identity_parts).encode("utf-8")).hexdigest()


def _effective_direction(source_direction: str, source_amount: Decimal) -> str | None:
    direction = source_direction.strip()
    if direction == "进账":
        effective = "credit"
    elif direction == "出账":
        effective = "debit"
    else:
        return None
    if source_amount < 0:
        return "debit" if effective == "credit" else "credit"
    return effective


def _parse_datetime(value: object) -> datetime:
    if _is_blank(value):
        raise _RequiredValueMissing()
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, str):
        normalized = value.strip().replace("年", "-").replace("月", "-").replace("日", "")
        return datetime.fromisoformat(normalized)
    raise ValueError("记账日期无法解析")


def _parse_direction(value: object) -> str:
    if _is_blank(value):
        raise _RequiredValueMissing()
    return str(value).strip()


def _parse_decimal(value: object) -> Decimal:
    if _is_blank(value):
        raise _RequiredValueMissing()
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return Decimal(str(value))
    normalized = str(value).strip().replace(",", "").replace("，", "")
    parentheses = re.fullmatch(r"\((.+)\)", normalized)
    if parentheses:
        normalized = f"-{parentheses.group(1)}"
    return Decimal(normalized)


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _best_effort_date(value: object) -> str | None:
    try:
        return _parse_datetime(value).date().isoformat()
    except (TypeError, ValueError, _RequiredValueMissing):
        return None


def _blocked_mapping_result() -> NormalizationResult:
    return NormalizationResult(
        code=CODE_MAPPING_FINGERPRINT_MISMATCH,
        message_zh="当前表头指纹与已确认映射不一致，已阻断逐笔处理；请回到R3重新确认字段。",
        input_transaction_rows=0,
        normalized_transaction_rows=0,
        transactions=(),
        invalid_rows=(),
        chosen_order=None,
        booking_datetime_regression_rows=(),
        reconciliation=ReconciliationSummary(0, 0, None, ()),
    )


class _RequiredValueMissing(ValueError):
    """Internal distinction between missing and unparsable required values."""
